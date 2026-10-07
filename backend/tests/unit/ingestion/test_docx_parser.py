from __future__ import annotations

import io

import pytest
from docx import Document

from policy_api.ingestion.parsers.base import ParseError
from policy_api.ingestion.parsers.docx import parse_docx


def document_bytes() -> bytes:
    document = Document()
    document.add_heading("休假制度", level=1)
    document.add_paragraph("适用于全体员工。")
    document.add_paragraph("提前申请", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "类型"
    table.cell(0, 1).text = "天数"
    buffer = io.BytesIO(); document.save(buffer); return buffer.getvalue()


def test_docx_parser_preserves_heading_list_and_table_order() -> None:
    blocks = parse_docx(document_bytes())
    assert [(b.block_type, b.text) for b in blocks] == [
        ("heading", "休假制度"), ("paragraph", "适用于全体员工。"),
        ("list_item", "提前申请"), ("table", "类型 | 天数"),
    ]
    assert blocks[1].heading_path == ("休假制度",)
    assert "类型" in blocks[3].text and "天数" in blocks[3].text


def test_docx_parser_rejects_corrupt_document() -> None:
    with pytest.raises(ParseError, match="docx_invalid"):
        parse_docx(b"not-a-docx")
